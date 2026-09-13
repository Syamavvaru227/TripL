"""Generate the SIH 2025 Idea Submission deck for TripL.

Faithfully reproduces the official SIH Idea Submission template:
  * white background, no title-slide footer
  * SIH blue (#006FBF) footer bar with "@SIH Idea submission- Template"
  * SIH bulb logo top-right, team badge top-left
  * centred Poppins-style section headings
  * lavender (#DADCEC) content pills with red (#FF3131) keyword highlights
  * medium-blue (#4A7EBB) commentary lines

All font sizes and positions are the template's own values scaled from its
20in x 11.25in master to a 13.333in x 7.5in 16:9 slide (factor 2/3).

Run:  python docs/generate_sih_ppt.py
Out:  TripL_SIH_Idea_Submission.pptx  (repository root)
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ------------------------------------------------- template palette --------
SIH_BLUE = RGBColor(0x00, 0x6F, 0xBF)      # footer bar / accents
BLUE_MED = RGBColor(0x4A, 0x7E, 0xBB)      # commentary lines & links
NAVY = RGBColor(0x1F, 0x49, 0x7D)          # title-slide heading, subtitle
PILL = RGBColor(0xDA, 0xDC, 0xEC)          # lavender content pill
PILL_LN = RGBColor(0xC3, 0xC6, 0xDD)
RED_HL = RGBColor(0xFF, 0x31, 0x31)        # template's keyword highlight
INK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x54, 0x54, 0x54)
LOGO_TXT = RGBColor(0x41, 0x58, 0x61)
WATERMARK = RGBColor(0xEC, 0xEC, 0xEC)
LINE = RGBColor(0xD2, 0xD6, 0xE0)
GRID = RGBColor(0xB9, 0xC6, 0xD8)
GREEN = RGBColor(0x19, 0x88, 0x3F)
ORANGE = RGBColor(0xEB, 0x86, 0x1A)

# The template uses Poppins (display) + Arimo (Arial) + Calibri.
# Century Gothic is Office-bundled and geometric like Poppins, so it is used
# as the portable stand-in. Install Poppins from Google Fonts for exact match.
DISPLAY = "Century Gothic"
SANS = "Arial"
BADGE = "Calibri"

SW, SH = 13.333, 7.5
FOOT_Y, FOOT_H = 6.947, 0.553
LEFT = 0.30


# ---------------------------------------------------------------- helpers --
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if s.has_text_frame:
        s.text_frame.clear()
        s.text_frame.word_wrap = True
    return s


def text(slide, x, y, w, h, content, size=12, color=INK, bold=False, italic=False,
         font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(content if isinstance(content, list) else [content]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(0)
        for r in (para if isinstance(para, list) else [para]):
            if isinstance(r, str):
                r = {"t": r}
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("font", font)
            f.size = Pt(r.get("size", size))
            f.bold = r.get("bold", bold)
            f.italic = r.get("italic", italic)
            f.color.rgb = r.get("color", color)
    return tb


def sih_logo(slide, x, y, scale=1.0):
    bw, bh = 0.30 * scale, 0.34 * scale
    rect(slide, x, y, bw, bh, fill=None, line=ORANGE, lw=1.4 * scale,
         shape=MSO_SHAPE.OVAL)
    rect(slide, x + 0.045 * scale, y + 0.115 * scale, 0.075 * scale, 0.012 * scale, fill=ORANGE)
    rect(slide, x + 0.045 * scale, y + 0.155 * scale, 0.105 * scale, 0.012 * scale, fill=ORANGE)
    rect(slide, x + 0.045 * scale, y + 0.195 * scale, 0.075 * scale, 0.012 * scale, fill=ORANGE)
    rect(slide, x + 0.055 * scale, y + 0.105 * scale, 0.012 * scale, 0.055 * scale, fill=ORANGE)
    rect(slide, x + 0.085 * scale, y + 0.145 * scale, 0.012 * scale, 0.075 * scale, fill=ORANGE)
    strip = rect(slide, x + 0.155 * scale, y + 0.075 * scale, 0.10 * scale,
                 0.23 * scale, fill=GREEN)
    tf = strip.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(("1010", "0110", "1001")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        r = p.add_run(); r.text = row
        r.font.name = SANS; r.font.size = Pt(3.0 * scale); r.font.bold = True
        r.font.color.rgb = WHITE
    rect(slide, x + 0.085 * scale, y + bh, 0.13 * scale, 0.022 * scale, fill=ORANGE)
    rect(slide, x + 0.095 * scale, y + bh + 0.028 * scale, 0.11 * scale, 0.018 * scale,
         fill=ORANGE)
    text(slide, x - 0.02 * scale, y + bh + 0.055 * scale, bw + 0.04 * scale,
         0.16 * scale, "SIH", size=8.5 * scale, color=LOGO_TXT, bold=True, font=SANS,
         align=PP_ALIGN.CENTER)
    text(slide, x + bw + 0.20 * scale, y + 0.01 * scale, 1.45 * scale, 0.62 * scale,
         [[{"t": "SMART INDIA"}], [{"t": "HACKATHON"}], [{"t": "2025"}]],
         size=7.6 * scale, color=LOGO_TXT, bold=True, font=SANS, spacing=1.05)


def team_badge(slide, name):
    e = rect(slide, LEFT, 0.45, 1.52, 0.50, fill=None,
             line=RGBColor(0x9A, 0xB2, 0xCE), lw=1.0, shape=MSO_SHAPE.OVAL)
    tf = e.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name
    r.font.name = BADGE; r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = INK


def footer(slide, num):
    rect(slide, 0, FOOT_Y, SW, FOOT_H, fill=SIH_BLUE)
    text(slide, 3.9, FOOT_Y, 6.0, FOOT_H, "@SIH Idea submission- Template",
         size=12, color=WHITE, font=SANS, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 12.25, FOOT_Y, 0.77, FOOT_H, str(num), size=12, color=WHITE, bold=True,
         font=SANS, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def heading(slide, title, y=0.30, size=23):
    text(slide, 1.5, y, SW - 3.0, 0.66, title, size=size, color=INK, bold=True,
         font=DISPLAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def pill_text(slide, x, y, w, h, runs, size=12, align=PP_ALIGN.CENTER,
              fill=PILL, radius=0.10):
    s = rect(slide, x, y, w, h, fill=fill, line=PILL_LN, lw=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.06
    for r in (runs if isinstance(runs, list) else [runs]):
        if isinstance(r, str):
            r = {"t": r}
        run = p.add_run(); run.text = r["t"]
        run.font.name = r.get("font", SANS)
        run.font.size = Pt(r.get("size", size))
        run.font.bold = r.get("bold", False)
        run.font.color.rgb = r.get("color", INK)
    return s


# ================================================================= slide 1 ===
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    rect(s, 4.35, 0.55, 7.40, 6.30, fill=WATERMARK, shape=MSO_SHAPE.HEXAGON)
    rect(s, 5.30, 1.55, 5.60, 4.70, fill=WHITE, shape=MSO_SHAPE.HEXAGON)

    sih_logo(s, 10.62, 0.28, scale=1.20)
    text(s, 0.85, 0.24, 10.40, 0.72, "SMART INDIA HACKATHON 2025", size=40,
         color=NAVY, bold=True, font="Times New Roman", align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)

    items = [
        (1.28, 0.52, "Problem Statement ID ", "\u2013"),
        (2.14, 1.00, "Problem Statement Title ", "\u2013 AI-Powered Intelligent Travel "
                                                 "Planning for Indian Tourism"),
        (3.43, 0.52, "Theme ", "\u2013 Travel & Tourism"),
        (4.29, 0.52, "PS Category ", "\u2013 Software"),
        (5.15, 0.52, "Team ID ", "\u2013"),
        (6.01, 0.52, "Team Name: ", "TripL"),
    ]
    for y, h, label, value in items:
        rect(s, 1.06, y + 0.19, 0.085, 0.085, fill=INK, shape=MSO_SHAPE.OVAL)
        text(s, 1.32, y, 9.4, h,
             [[{"t": label, "bold": True}, {"t": value, "bold": True}]],
             size=26, color=INK, font=SANS, spacing=1.06)
    return s


# ================================================================= slide 2 ===
def slide_solution(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    team_badge(s, "TripL")
    sih_logo(s, 10.72, 0.28, scale=1.15)
    text(s, 2.55, 0.22, 8.10, 0.82, "TripL \u2014 AI-Powered Intelligent Travel "
                                    "Planning and Tourist Discovery",
         size=22, color=INK, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.08)
    text(s, LEFT, 1.22, 9.20, 0.42,
         "Proposed Solution (Describe your Idea/Solution/Prototype)",
         size=20, color=NAVY, font=DISPLAY)

    body = [
        [{"t": "Discovers real tourist places within 30 km of "},
         {"t": "ANY Indian location", "color": RED_HL, "bold": True},
         {"t": " using live Wikipedia GeoSearch and OpenStreetMap data \u2014 no static "
               "database, no city limits."}],
        [{"t": "Filters out villages, railway stations and irrelevant entries with "},
         {"t": "multi-layer keyword, category and confidence checks", "color": RED_HL,
          "bold": True},
         {"t": " so only genuine tourist places reach the user."}],
        [{"t": "Compares car, bike, bus, auto and walking on real road distances, travel "
               "time and rupee cost using the "},
         {"t": "OSRM routing engine", "color": RED_HL, "bold": True}, {"t": "."}],
        [{"t": "Generates a "},
         {"t": "personalised AI itinerary", "color": RED_HL, "bold": True},
         {"t": " from budget, available time, interests, place-type filters and transport "
               "preference \u2014 with a reason for every stop."}],
        [{"t": "Adds a "},
         {"t": "Know-India cultural layer", "color": RED_HL, "bold": True},
         {"t": " (history, festivals, cuisine, etiquette) and responsible-tourism badges, "
               "plus one-tap booking with Uber, Ola and Rapido."}],
    ]
    cy, ph, pgap = 1.74, 0.80, 0.15
    for runs in body:
        pill_text(s, 0.42, cy, 7.55, ph, runs, size=12)
        cy += ph + pgap

    # radial capability diagram
    cx, cy0, d = 9.98, 3.35, 1.15
    rect(s, cx, cy0, d, d, fill=SIH_BLUE, shape=MSO_SHAPE.OVAL)
    text(s, cx, cy0, d, d, [[{"t": "TripL", "size": 12, "bold": True, "color": WHITE,
                              "font": DISPLAY}],
                            [{"t": "AI Engine", "size": 8.5, "bold": True,
                              "color": WHITE}]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    left_nodes = ["Live Place\nDiscovery", "Smart Noise\nFiltering", "Real-Road\nRouting"]
    right_nodes = ["AI Itinerary\nEngine", "Know India\nCulture", "Responsible\nTourism"]
    nh, ngap, ny0 = 0.78, 0.36, 2.40
    for i, lab in enumerate(left_nodes):
        y = ny0 + i * (nh + ngap)
        pill_text(s, 8.10, y, 1.60, nh, lab, size=8.8, fill=PILL, radius=0.16)
        rect(s, 9.70, y + nh / 2 - 0.012, cx - 9.70, 0.024, fill=LINE)
    for i, lab in enumerate(right_nodes):
        y = ny0 + i * (nh + ngap)
        pill_text(s, 11.45, y, 1.55, nh, lab, size=8.8, fill=PILL, radius=0.16)
        rect(s, cx + d, y + nh / 2 - 0.012, 11.45 - (cx + d), 0.024, fill=LINE)

    text(s, 0.42, 6.36, 12.60, 0.5,
         "Search anywhere  \u2192  discover what is around it  \u2192  compare how to get "
         "there  \u2192  get an AI plan  \u2192  understand its story  \u2192  travel "
         "responsibly",
         size=11.5, color=BLUE_MED, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
    footer(s, 2)
    return s


# ================================================================= slide 3 ===
def slide_technical(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    team_badge(s, "TripL")
    sih_logo(s, 10.72, 0.28, scale=1.15)
    heading(s, "TECHNICAL APPROACH", y=0.34)

    text(s, 0.22, 1.16, 3.0, 0.28, "Techstack Used:", size=12, color=INK, bold=True,
         font=DISPLAY)
    stack = ["1. React 18 + Vite", "2. Python (FastAPI)", "3. Tailwind CSS",
             "4. Leaflet + OpenStreetMap", "5. SQLAlchemy + MySQL", "6. JWT Auth",
             "7. Wikipedia + Nominatim", "8. OSRM Routing Engine"]
    text(s, 0.22, 1.48, 3.1, 1.9, [[{"t": t}] for t in stack], size=10.5,
         color=INK, spacing=1.30)

    text(s, 0.22, 3.62, 3.0, 0.26, "Prototype:", size=12, color=INK, bold=True,
         font=DISPLAY)
    text(s, 0.22, 3.90, 3.1, 0.26,
         [[{"t": "Link  ", "bold": True, "size": 10.5},
           {"t": "TripL Web App", "color": BLUE_MED, "bold": True, "size": 10.5}]])
    text(s, 0.22, 4.28, 3.2, 0.26, "Prototype Video Link:", size=12, color=INK,
         bold=True, font=DISPLAY)
    text(s, 0.22, 4.56, 3.1, 0.26,
         [[{"t": "Youtube Video", "color": BLUE_MED, "bold": True, "size": 10.5}]])

    steps = [
        ("User Input", "City search or\nlive GPS location"),
        ("Geocoding", "Nominatim resolves\nlatitude / longitude"),
        ("Live Discovery", "Wikipedia GeoSearch\n+ TextSearch"),
        ("Smart Filtering", "Reject noise, de-dupe,\ninfer category"),
        ("Enrichment", "Timings, crowds, images,\nKnow-India content"),
        ("Routing", "OSRM real road\ndistance and duration"),
        ("AI Planner", "Weighted scoring and\nschedule optimisation"),
        ("Delivery", "Timeline, map, cost\nsummary, save and share"),
    ]
    x0, y0, bw, bh, gap = 3.52, 1.20, 2.18, 1.12, 0.20
    for i, (name, desc) in enumerate(steps):
        row, col = divmod(i, 4)
        x = x0 + col * (bw + gap)
        y = y0 + row * 1.68
        rect(s, x, y, bw, bh, fill=WHITE, line=GRID, lw=0.9,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        rect(s, x, y, bw, 0.07, fill=SIH_BLUE if row == 0 else BLUE_MED)
        text(s, x + 0.10, y + 0.16, bw - 0.20, 0.24, name, size=10, color=SIH_BLUE,
             bold=True, font=DISPLAY)
        text(s, x + 0.10, y + 0.43, bw - 0.20, 0.60, desc, size=8.8, color=MUTED,
             spacing=1.05)
        if col < 3:
            rect(s, x + bw + 0.035, y + 0.49, 0.13, 0.13, fill=SIH_BLUE,
                 shape=MSO_SHAPE.ISOSCELES_TRIANGLE).rotation = 90

    pill_text(s, 3.52, 4.62, 9.50, 0.80,
              [{"t": "Results found?   ", "bold": True, "color": SIH_BLUE, "size": 10.5},
               {"t": "Yes \u2192 score, schedule and cost the journey.      ", "size": 10.3},
               {"t": "No \u2192 ", "bold": True, "color": SIH_BLUE, "size": 10.5},
               {"t": "widen the radius, relax filters or fall back to the nearest city "
                     "centre.", "size": 10.3}],
              align=PP_ALIGN.CENTER)
    rect(s, 3.52, 5.56, 9.50, 0.80, fill=WHITE, line=GRID, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    text(s, 3.68, 5.56, 9.18, 0.80,
         [[{"t": "Continuous improvement:   ", "bold": True, "color": SIH_BLUE,
            "size": 10.5},
           {"t": "response caching, async httpx batching and user feedback on visited "
                 "places steadily improve place ranking and itinerary quality.",
            "size": 10.3}]],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    footer(s, 3)
    return s


# ================================================================= slide 4 ===
def slide_feasibility(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    team_badge(s, "TripL")
    sih_logo(s, 10.72, 0.28, scale=1.15)
    heading(s, "FEASIBILITY AND VIABILITY", y=0.30)

    rows = [
        ("Technical \u2013 Achievable",
         "Latency & real-time performance across third-party APIs.",
         "Fully async httpx calls, in-memory caching and containerised FastAPI keep "
         "discovery and itinerary generation responsive."),
        ("Accuracy \u2013 Effective",
         "Irrelevant results and noise: villages, railway stations, schools, dish names.",
         "Layered keyword rejection plus category inference and confidence scoring remove "
         "non-tourist entries before planning."),
        ("Data \u2013 Enrichable, Compliant",
         "Keeping place data fresh without maintaining a dataset.",
         "Live Wikipedia, Wikidata, Nominatim and OpenStreetMap feeds \u2014 no licence "
         "cost, no manual updates, always current."),
        ("Cost \u2013 Scalable",
         "Server and API expenses when scaled to a national audience.",
         "Zero-licence data stack, free-tier deployment and horizontal scaling; OSRM "
         "falls back to Haversine when unreachable."),
        ("Adoption \u2013 Practical",
         "Trust and real usage by tourists, hosts and tourism boards.",
         "Free mobile-first web app, no install, shareable itineraries and local-language "
         "support in the roadmap."),
    ]

    tx, ty = 0.22, 1.06
    col = [2.55, 4.35, 5.90]
    rect(s, tx, ty, sum(col), 0.50, fill=PILL, line=PILL_LN, lw=0.75)
    cx = tx
    for i, h in enumerate(["Analysis of the feasibility of the idea",
                           "Potential challenges and risks",
                           "Strategies for overcoming these challenges"]):
        text(s, cx + 0.10, ty, col[i] - 0.20, 0.50, h, size=11, color=INK, bold=True,
             font=DISPLAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
        cx += col[i]
    y = ty + 0.50
    for i, r in enumerate(rows):
        h = 0.64
        rect(s, tx, y, sum(col), h, fill=WHITE, line=LINE, lw=0.6)
        cx = tx
        for j, v in enumerate(r):
            text(s, cx + 0.11, y + 0.05, col[j] - 0.22, h - 0.10, v,
                 size=10.2 if j else 11, color=INK, bold=(j == 0),
                 font=DISPLAY if j == 0 else SANS, spacing=1.04,
                 anchor=MSO_ANCHOR.MIDDLE)
            cx += col[j]
        y += h + 0.05

    text(s, 0.22, y + 0.16, 12.80, 0.30,
         "Given that feasibility and viability being our top priority we will ensure to "
         "keep everything", size=15, color=BLUE_MED, bold=True, font=DISPLAY,
         align=PP_ALIGN.CENTER)
    lines = [
        "Technically viable \u2013 built on proven open tools (React, FastAPI, Wikipedia, OSM), no heavy on-device ML.",
        "Economically feasible \u2013 zero data-licensing cost, free-tier cloud setup, scales from a district to national level.",
        "Operationally practical \u2013 simple web app, live discovery, automatic planning, no manual data maintenance.",
    ]
    text(s, 0.42, y + 0.54, 12.40, 1.05, [[{"t": l}] for l in lines], size=12,
         color=BLUE_MED, spacing=1.22)
    footer(s, 4)
    return s


# ================================================================= slide 5 ===
def slide_impact(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    team_badge(s, "TripL")
    sih_logo(s, 10.72, 0.28, scale=1.15)
    heading(s, "IMPACT AND BENEFITS", y=0.30)

    for x, w, title in ((0.24, 6.25, "Potential Impact on Target Audience"),
                        (6.80, 6.30, "Benefits of the solution")):
        rect(s, x, 1.28, w, 0.46, fill=PILL, line=PILL_LN, lw=0.75)
        text(s, x, 1.28, w, 0.46, title, size=12.5, color=INK, bold=True, font=DISPLAY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    left = [
        ("Tourists & Families", "discover hidden gems within 30 km and plan a complete day "
                                "in minutes instead of hours."),
        ("Local Communities & Artisans", "visibility for homestays, handicrafts and street "
                                         "food pushes income to the grassroots."),
        ("Tier-2 & Tier-3 Towns", "destinations outside the tourist circuit finally get a "
                                  "digital presence."),
        ("State Tourism Boards", "data-backed insight into under-visited places helps spread "
                                 "footfall and ease crowding."),
        ("Budget & Student Travellers", "transparent mode-wise costing makes travel "
                                        "affordable and predictable."),
    ]
    right = [
        ("Social Benefits", "makes India's heritage accessible to everyone and promotes "
                            "inclusive, respectful tourism."),
        ("Economic Benefits", "boosts micro-economies \u2014 guides, homestays, drivers, "
                              "artisans and local eateries."),
        ("Environmental Benefits", "highlights eco-friendly spots and public-transport "
                                   "options to cut the carbon footprint of every trip."),
        ("Technological Benefits", "a reusable, API-driven discovery engine that any state "
                                   "tourism portal can adopt."),
        ("Strategic Benefits", "supports Dekho Apna Desh and Digital India with fully "
                               "homegrown, open technology."),
    ]

    def numbered(entries, x, w):
        cy = 1.90
        for i, (lead, rest) in enumerate(entries, start=1):
            text(s, x + 0.04, cy, 0.30, 0.26, f"{i}.", size=12, color=INK, bold=True,
                 font=DISPLAY)
            text(s, x + 0.38, cy, w - 0.42, 0.78,
                 [[{"t": lead + " \u2013 ", "bold": True, "size": 11.5},
                   {"t": rest, "size": 11.5}]], spacing=1.08)
            cy += 0.60

    numbered(left, 0.24, 6.25)
    numbered(right, 6.80, 6.30)

    rect(s, 9.60, 5.22, 2.50, 0.72, fill=WHITE, line=SIH_BLUE, lw=1.4,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    text(s, 9.60, 5.22, 2.50, 0.72, [[{"t": "Plan in minutes."}],
                                     [{"t": "Travel responsibly."}],
                                     [{"t": "Support local."}]],
         size=10, color=BLUE_MED, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.02)
    footer(s, 5)
    return s


# ================================================================= slide 6 ===
def slide_research(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    team_badge(s, "TripL")
    sih_logo(s, 10.72, 0.28, scale=1.15)
    text(s, 0.22, 1.24, 2.6, 0.26, "Research Survey", size=12, color=INK, bold=True,
         font=DISPLAY)
    heading(s, "RESEARCH AND REFERENCES", y=0.30)

    rows = [
        ("Discovery", "Static curated DB; only major cities covered",
         "Live discovery for ANY Indian location within 30 km"),
        ("Data", "Manually maintained, goes stale quickly",
         "Real-time Wikipedia + OpenStreetMap, always current"),
        ("Planning", "Fixed packages and static templates",
         "AI itinerary from live budget, time, interests and place types"),
        ("Cost", "Vague estimates, no mode-wise comparison",
         "Real road distance and mode-wise cost via OSRM"),
        ("Culture", "A one-line description at best",
         "Know-India layer: history, festivals, cuisine and etiquette"),
        ("Sustainability", "No responsible-tourism signal",
         "Eco, community and artisan badges on every destination"),
    ]

    tx, ty = 0.22, 1.56
    col = [1.90, 5.05, 5.85]
    rect(s, tx, ty, sum(col), 0.40, fill=SIH_BLUE)
    cx = tx
    for i, h in enumerate(["Aspect", "Conventional Systems", "Proposed System"]):
        text(s, cx + 0.12, ty, col[i] - 0.24, 0.40, h, size=12, color=WHITE, bold=True,
             font=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        cx += col[i]
    y = ty + 0.40
    for i, r in enumerate(rows):
        h = 0.34
        rect(s, tx, y, sum(col), h, fill=WHITE if i % 2 else RGBColor(0xF4, 0xF6, 0xFB),
             line=LINE, lw=0.5)
        cx = tx
        for j, v in enumerate(r):
            text(s, cx + 0.12, y, col[j] - 0.24, h, v, size=10,
                 bold=(j == 0),
                 color=INK if j == 0 else (BLUE_MED if j == 2 else MUTED),
                 font=DISPLAY if j == 0 else SANS, anchor=MSO_ANCHOR.MIDDLE)
            cx += col[j]
        y += h

    text(s, 0.22, y + 0.20, 4.0, 0.26, "Research Paper:", size=11.5, color=INK,
         bold=True, font=DISPLAY)
    refs = [
        "[1] Wikimedia Foundation, \u201cMediaWiki API: GeoSearch and TextSearch\u201d "
        "[Online]. Available: https://www.mediawiki.org/wiki/API:Geosearch",
        "[2] OpenStreetMap Contributors, \u201cNominatim Geocoding and Reverse Geocoding "
        "API\u201d [Online]. Available: https://nominatim.org/release-docs/latest/api/",
        "[3] Project OSRM, \u201cOpen Source Routing Machine: road distance and "
        "duration\u201d [Online]. Available: https://project-osrm.org/",
    ]
    cy = y + 0.48
    for r in refs:
        text(s, 0.26, cy, 12.60, 0.30, r, size=10, color=INK, spacing=1.06)
        cy += 0.32

    text(s, 0.22, cy + 0.10, 4.0, 0.26, "Dataset:", size=11.5, color=INK, bold=True,
         font=DISPLAY)
    data = ["[1] Wikipedia GeoSearch \u2013 geo-tagged tourist articles",
            "[2] Wikidata \u2013 descriptions, history and cultural context",
            "[3] OpenStreetMap POI tags \u2013 categories, coordinates, hours",
            "[4] OSRM India road graph \u2013 real driving distance and time"]
    cy += 0.40
    for i, d in enumerate(data):
        text(s, 0.26 + (i % 2) * 6.55, cy + (i // 2) * 0.30, 6.30, 0.28, d, size=10,
             color=INK)
    footer(s, 6)
    return s


# ==================================================================== main ===
def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    slide_title(prs)
    slide_solution(prs)
    slide_technical(prs)
    slide_feasibility(prs)
    slide_impact(prs)
    slide_research(prs)

    out = Path(__file__).resolve().parent.parent / "TripL_SIH_Idea_Submission.pptx"
    prs.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    main()
